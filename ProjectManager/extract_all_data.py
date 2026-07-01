import openpyxl
import json
import datetime
import sys

# Reconfigure stdout to use UTF-8
sys.stdout.reconfigure(encoding='utf-8')

def clean_val(val):
    if val is None:
        return None
    if isinstance(val, datetime.datetime):
        return val.strftime("%Y-%m-%d")
    return val

def clean_date_str(val):
    if val is None:
        return None
    if isinstance(val, datetime.datetime):
        return val.strftime("%Y-%m-%d")
    val_str = str(val).strip()
    # If it is like "15/12/2025" convert to "2025-12-15"
    if "/" in val_str:
        parts = val_str.split("/")
        if len(parts) == 3:
            # Check if year is first or last
            if len(parts[2]) == 4:
                return f"{parts[2]}-{parts[1].zfill(2)}-{parts[0].zfill(2)}"
            elif len(parts[0]) == 4:
                return f"{parts[0]}-{parts[1].zfill(2)}-{parts[2].zfill(2)}"
    return val_str

print("Loading workbook...")
wb = openpyxl.load_workbook('c:/Users/luise/OneDrive/Desarrollo/Manager/ProjectManager/Plan de Trabajo MesaGo.xlsx', data_only=True)

# 1. Extract Activities
print("Extracting activities...")
act_sheet = wb['Actividades']
activities = []
# Row 1 has headers: Sección, Descripción, Estatus, Etiqueta, Fecha estimada original, Fecha real, Días, Fecha estimada
for r in range(2, 200): # Scan up to 200 rows
    section = clean_val(act_sheet.cell(r, 1).value)
    description = clean_val(act_sheet.cell(r, 2).value)
    status = clean_val(act_sheet.cell(r, 3).value)
    tag = clean_val(act_sheet.cell(r, 4).value)
    orig_est_date = clean_date_str(act_sheet.cell(r, 5).value)
    real_date = clean_date_str(act_sheet.cell(r, 6).value)
    days = act_sheet.cell(r, 7).value
    est_date = clean_date_str(act_sheet.cell(r, 8).value)
    
    if not section and not description:
        # Check if we reached the end
        if all(act_sheet.cell(r + i, 1).value is None for i in range(1, 10)):
            break
        continue
        
    activities.append({
        "section": section,
        "description": description,
        "status": status,
        "tag": tag,
        "original_estimated_date": orig_est_date,
        "real_date": real_date,
        "days": int(days) if days is not None and str(days).isdigit() else None,
        "estimated_date": est_date
    })

# 2. Extract Components, MVP Priorities and Timeline from Dashboard
print("Extracting dashboard data...")
db_sheet = wb['Dashboard']

components = []
# Components are on rows 17 to 25
for r in range(17, 26):
    name = clean_val(db_sheet.cell(r, 1).value)
    weight = db_sheet.cell(r, 2).value
    progress = db_sheet.cell(r, 3).value
    if name:
        components.append({
            "name": name,
            "weight": float(weight) if weight is not None else 0.0,
            "progress": float(progress) if progress is not None else 0.0
        })

mvp_priorities = []
# MVP priorities are on rows 21 to 27, columns 7 to 11
for r in range(21, 28):
    num = clean_val(db_sheet.cell(r, 7).value)
    priority = clean_val(db_sheet.cell(r, 8).value)
    func = clean_val(db_sheet.cell(r, 9).value)
    criticity = clean_val(db_sheet.cell(r, 10).value)
    status = clean_val(db_sheet.cell(r, 11).value)
    if func:
        mvp_priorities.append({
            "priority": priority or "CRÍTICO",
            "functionality": func,
            "criticity": criticity or "⭐⭐⭐",
            "status": status or "Pendiente"
        })

timeline = []
# Executive timeline logs are on rows 6 to 13, columns 7 to 11
for r in range(6, 14):
    date = clean_date_str(db_sheet.cell(r, 7).value)
    comp = clean_val(db_sheet.cell(r, 8).value)
    task = clean_val(db_sheet.cell(r, 9).value)
    pct = db_sheet.cell(r, 10).value
    notes = clean_val(db_sheet.cell(r, 11).value)
    if date and task:
        timeline.append({
            "date": date,
            "component": comp or "Desconocido",
            "task_description": task,
            "percentage_logrado": float(pct) if pct is not None else 1.0,
            "notes": notes
        })

data_export = {
    "activities": activities,
    "components": components,
    "mvp_priorities": mvp_priorities,
    "timeline": timeline
}

print(f"Extracted {len(activities)} activities.")
print(f"Extracted {len(components)} components.")
print(f"Extracted {len(mvp_priorities)} MVP priorities.")
print(f"Extracted {len(timeline)} timeline logs.")

with open('c:/Users/luise/OneDrive/Desarrollo/Manager/ProjectManager/data_export.json', 'w', encoding='utf-8') as f:
    json.dump(data_export, f, indent=2, ensure_ascii=False)

print("Data exported successfully to data_export.json")
